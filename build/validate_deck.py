import argparse
import json
import pathlib
from pptx import Presentation

BASE = pathlib.Path(__file__).resolve().parents[1]

ap = argparse.ArgumentParser()
ap.add_argument("--pptx", default=str(BASE / "output" / "integration_farmer_proposal.pptx"))
ap.add_argument("--outline", default=str(BASE / "content" / "outline.json"))
args = ap.parse_args()

PPTX = pathlib.Path(args.pptx)
OUTLINE = json.loads(pathlib.Path(args.outline).read_text(encoding="utf-8"))

prs = Presentation(str(PPTX))
errors = []
slides = list(prs.slides)
expected = len(OUTLINE["slides"])

if len(slides) != expected:
    errors.append(f"slide count {len(slides)} != outline {expected}")


def expected_pictures(d):
    if d["layout"] == "pillars":
        return len(d["cards"])
    if d["layout"] == "mechanism2":
        steps = sum(len(st["photos"]) for st in d["steps"])
        return steps + 1
    return 1 if "image" in d else 0


total_pictures = 0
for idx, (slide, d) in enumerate(zip(slides, OUTLINE["slides"]), start=1):
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    total_pictures += len(pics)
    exp = expected_pictures(d)
    if len(pics) != exp:
        errors.append(f"slide {idx}: pictures {len(pics)} != expected {exp}")
    if not any(sh.has_text_frame and sh.text_frame.text.strip() for sh in slide.shapes):
        errors.append(f"slide {idx}: no text content")
    if idx > 1:
        has_notes = slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
        if not has_notes:
            errors.append(f"slide {idx}: missing speaker notes")

tables = [sh for s in slides for sh in s.shapes if sh.has_table]
if len(tables) != 1:
    errors.append(f"tables found {len(tables)} != 1")
else:
    t = tables[0].table
    if len(t.rows) != 8 or len(t.columns) != 4:
        errors.append(f"table size {len(t.rows)}x{len(t.columns)} != 8x4")

expected_total = sum(expected_pictures(d) for d in OUTLINE["slides"])
if total_pictures != expected_total:
    errors.append(f"total pictures {total_pictures} != {expected_total}")

print(f"slides={len(slides)} pictures={total_pictures} tables={len(tables)}")
if errors:
    print("VALIDATION ERRORS:")
    for e in errors:
        print(" -", e)
    raise SystemExit(1)
print("Structure validation PASSED")
