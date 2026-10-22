import json
import pathlib
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn, nsdecls
from PIL import Image

BASE = pathlib.Path(__file__).resolve().parents[1]
OUTLINE_PATH = BASE / "content" / "outline.json"
OUT_PATH = BASE / "output" / "integration_farmer_proposal.pptx"
THEME = json.loads((BASE / "content" / "theme.json").read_text(encoding="utf-8"))


def load_outline(path=None):
    global OUTLINE, ASSETS
    OUTLINE = json.loads(pathlib.Path(path or OUTLINE_PATH).read_text(encoding="utf-8"))
    ASSETS = {k: BASE / v for k, v in OUTLINE["assets"].items()}
    return OUTLINE


OUTLINE = None
ASSETS = None
load_outline()

COL = {k: RGBColor.from_string(v) for k, v in THEME["colors"].items()}
SZ = THEME["sizes"]
LAY = THEME["layout"]
FONT_H = THEME["fonts"]["heading"]
FONT_B = THEME["fonts"]["body"]
SW = THEME["slide"]["width_in"]
SH = THEME["slide"]["height_in"]
FOOTER_TEXT = "Proposal Program Integration Farmer  |  RT/RW"


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()
    return tf


def style_run(run, size, color, bold=False, font=None):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font or FONT_B


def txt(slide, x, y, w, h, text, size=None, color=None, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=None,
        line_spacing=1.05, space_after=0):
    tf = add_box(slide, x, y, w, h)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    style_run(r, size or SZ["body"], color or COL["text"], bold=bold, font=font)
    return tf


def add_bullets(slide, x, y, w, h, items, size=None, color=None,
                marker=None, bold=False, line_spacing=1.08, space_after=7,
                anchor=MSO_ANCHOR.TOP):
    tf = add_box(slide, x, y, w, h)
    tf.vertical_anchor = anchor
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        if marker:
            rm = p.add_run()
            rm.text = "\u2022  "
            style_run(rm, size or SZ["bullet"], marker, bold=True)
        r = p.add_run()
        r.text = item
        style_run(r, size or SZ["bullet"], color or COL["text"], bold=bold)
    return tf


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.clear()
    return sp


def card(slide, x, y, w, h, fill=None, border=True, radius=0.06):
    return shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h,
        fill=fill or COL["bg"],
        line=COL["card_border"] if border else None,
        line_w=1.2, radius=radius,
    )


def pic_fit(slide, key, x, y, w, h):
    path = ASSETS[key]
    iw, ih = Image.open(path).size
    box_ar = w / h
    img_ar = iw / ih
    if img_ar >= box_ar:
        pw, ph = w, w / img_ar
    else:
        ph, pw = h, h * img_ar
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def _round_pic(pic, radius=0.08):
    try:
        spPr = pic._element.spPr
        geom = parse_xml(
            '<a:prstGeom %s prst="roundRect"><a:avLst>'
            '<a:gd name="adj" fmla="val %d"/>'
            "</a:avLst></a:prstGeom>" % (nsdecls("a"), int(radius * 100000))
        )
        old = spPr.find(qn("a:prstGeom"))
        if old is not None:
            spPr.replace(old, geom)
        else:
            xfrm = spPr.find(qn("a:xfrm"))
            if xfrm is not None:
                xfrm.addnext(geom)
            else:
                spPr.insert(0, geom)
    except Exception as e:
        print(f"round_pic skipped: {e}")


def pic_cover(slide, key, x, y, w, h, radius=0.06):
    path = ASSETS[key]
    im = Image.open(path).convert("RGB")
    iw, ih = im.size
    target_ar = w / h
    img_ar = iw / ih
    if img_ar > target_ar:
        new_w = int(ih * target_ar)
        left = (iw - new_w) // 2
        box = (left, 0, left + new_w, ih)
    else:
        new_h = int(iw / target_ar)
        top = (ih - new_h) // 2
        box = (0, top, iw, top + new_h)
    outdir = BASE / "build" / "cropped"
    outdir.mkdir(parents=True, exist_ok=True)
    fname = outdir / f"{path.stem}_{int(w * 100)}x{int(h * 100)}.jpg"
    im.crop(box).save(fname, quality=92)
    pic = slide.shapes.add_picture(str(fname), Inches(x), Inches(y), Inches(w), Inches(h))
    _round_pic(pic, radius)
    return pic


def header(slide, headline, kicker=None, subhead=None, rule=True):
    if kicker:
        txt(slide, LAY["margin_in"], LAY["header_kicker_y_in"], 8.0, 0.3,
            kicker, size=SZ["kicker"], color=COL["primary"], bold=True)
    hy = LAY["header_headline_y_in"]
    if kicker:
        hy += 0.1
    txt(slide, LAY["margin_in"], hy, SW - 2 * LAY["margin_in"], 0.62,
        headline, size=SZ["headline"], color=COL["primary_dark"], bold=True,
        font=FONT_H)
    ry = LAY["header_rule_y_in"]
    if subhead:
        txt(slide, LAY["margin_in"], LAY["header_rule_y_in"] - 0.02, 11.0, 0.4,
            subhead, size=SZ["subhead"], color=COL["text_muted"])
        ry += 0.42
    if rule:
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              LAY["margin_in"], ry, 1.5, 0.07, fill=COL["secondary"], radius=0.5)
    return ry + 0.07


def footer(slide, idx):
    txt(slide, LAY["margin_in"], LAY["footer_y_in"], 7.0, 0.3,
        FOOTER_TEXT, size=SZ["footer"], color=COL["text_muted"])
    txt(slide, SW - LAY["margin_in"] - 1.2, LAY["footer_y_in"], 1.2, 0.3,
        str(idx), size=SZ["footer"], color=COL["text_muted"], align=PP_ALIGN.RIGHT)


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def chip(slide, x, y, w, h, text, size, fill, color, bold=True, align=PP_ALIGN.CENTER):
    sp = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, radius=0.5)
    tf = sp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    style_run(r, size, color, bold=bold)
    return sp


# ---------------------------------------------------------------- layouts

def layout_cover(prs, d):
    s = blank_slide(prs)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=COL["soft_bg"])
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, SW - 0.35, 0, 0.35, SH,
          fill=COL["primary"], radius=0.0)
    chip(s, 0.9, 1.0, 4.35, 0.44, d["kicker"], SZ["kicker"] + 0.5,
         COL["primary_light"], COL["primary"])
    txt(s, 0.9, 1.68, 6.2, 1.1, d["headline"], size=SZ["cover_title"],
        color=COL["primary_dark"], bold=True, font=FONT_H)
    txt(s, 0.9, 2.92, 6.2, 0.9, d["tagline"], size=SZ["cover_tagline"],
        color=COL["primary"], bold=True, font=FONT_H)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 3.72, 1.4, 0.07,
          fill=COL["accent"], radius=0.5)
    txt(s, 0.9, 4.0, 5.9, 1.3, d["subtitle"], size=SZ["cover_subtitle"],
        color=COL["text_muted"], line_spacing=1.25)
    txt(s, 0.9, 5.9, 5.9, 0.4, "Proposal Program Komunitas  |  Oktober 2026",
        size=13, color=COL["text"], bold=True)
    txt(s, 0.9, 6.35, 5.9, 0.4, OUTLINE["meta"]["approval_note"],
        size=12, color=COL["text_muted"])
    pic_cover(s, d["image"], 7.2, 1.35, 5.5, 5.0, radius=0.09)
    notes(s, d["notes"])
    return s


def layout_exec_summary(prs, d):
    s = blank_slide(prs)
    header(s, d["headline"])
    n = len(d["cards"])
    gap = LAY["gap_in"]
    m = LAY["margin_in"]
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 1.75, 2.45
    for i, c in enumerate(d["cards"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch, fill=COL["soft_bg"])
        txt(s, x + 0.25, cy + 0.3, cw - 0.5, 0.7, c["stat"], size=34,
            color=COL["primary"], bold=True, align=PP_ALIGN.CENTER, font=FONT_H)
        txt(s, x + 0.28, cy + 1.15, cw - 0.56, 1.1, c["label"], size=SZ["stat_label"],
            color=COL["text_muted"], align=PP_ALIGN.CENTER, line_spacing=1.15)
    by = 4.55
    for b in d["bullets"]:
        chip(s, m, by, SW - 2 * m, 0.62, b, 15.5, COL["primary_light"],
             COL["primary_dark"], align=PP_ALIGN.LEFT)
        by += 0.78
    footer(s, 2)
    notes(s, d["notes"])
    return s


def layout_content_stats(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    add_bullets(s, LAY["margin_in"], 1.85, 7.1, 4.6, d["bullets"],
                size=16.5, marker=COL["secondary"], space_after=12)
    rx, rw = 8.15, 4.45
    card(s, rx, 1.75, rw, 2.75, fill=COL["soft_bg"])
    pic_cover(s, d["image"], rx + 0.18, 1.93, rw - 0.36, 2.4, radius=0.07)
    sy = 4.7
    for st in d.get("stat_strip", []):
        card(s, rx, sy, rw, 1.0, fill=COL["primary_light"], border=False)
        txt(s, rx + 0.25, sy + 0.12, rw - 0.5, 0.42, st["stat"], size=20,
            color=COL["primary"], bold=True, font=FONT_H)
        txt(s, rx + 0.25, sy + 0.54, rw - 0.5, 0.4, st["label"], size=11.5,
            color=COL["text_muted"], line_spacing=1.0)
        sy += 1.15
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_content_image(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"], subhead=d["subhead"])
    add_bullets(s, LAY["margin_in"], 2.25, 6.5, 4.3, d["bullets"],
                size=16.5, marker=COL["secondary"], space_after=12)
    card(s, 7.55, 2.05, 5.1, 4.55, fill=COL["soft_bg"])
    pic_cover(s, d["image"], 7.75, 2.25, 4.7, 4.15, radius=0.07)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_pillars(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    n = len(d["cards"])
    gap = LAY["gap_in"]
    m = LAY["margin_in"]
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 1.7, 5.0
    for i, c in enumerate(d["cards"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch, fill=COL["bg"])
        pic_cover(s, c["image"], x + 0.18, cy + 0.18, cw - 0.36, 1.78, radius=0.07)
        txt(s, x + 0.2, cy + 2.12, cw - 0.4, 0.45, c["title"], size=20,
            color=COL["primary_dark"], bold=True, align=PP_ALIGN.CENTER, font=FONT_H)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + (cw - 0.9) / 2, cy + 2.64,
              0.9, 0.055, fill=COL["accent"], radius=0.5)
        add_bullets(s, x + 0.35, cy + 2.9, cw - 0.65, 1.9, c["bullets"],
                    size=SZ["card_bullet"], marker=COL["secondary"], space_after=7)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_diagram(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    pic_fit(s, d["image"], 1.1, 1.65, 11.1, 4.05)
    n = len(d["captions"])
    m = 0.7
    gap = 0.25
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy = 5.85
    for i, cap in enumerate(d["captions"]):
        x = m + i * (cw + gap)
        chip(s, x, cy, cw, 0.85, cap, SZ["caption"], COL["soft_bg"],
             COL["primary_dark"], bold=False)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_photo_captions(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    pic_cover(s, d["image"], 1.1, 1.65, 11.1, 4.05, radius=0.05)
    n = len(d["captions"])
    m = 0.7
    gap = 0.25
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy = 5.85
    for i, cap in enumerate(d["captions"]):
        x = m + i * (cw + gap)
        chip(s, x, cy, cw, 0.85, cap, SZ["caption"], COL["soft_bg"],
             COL["primary_dark"], bold=False)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_two_column(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    cy, ch = 1.7, 4.35
    x1, cw = 0.5, 4.33
    x2 = 5.13
    x3 = 8.43
    for x, title, bullets, band in (
        (x1, d["left_title"], d["left_bullets"], COL["primary"]),
        (x3, d["right_title"], d["right_bullets"], COL["accent"]),
    ):
        card(s, x, cy, cw, ch)
        band_sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, cy, cw, 0.62,
                        fill=band, radius=0.12)
        tf = band_sp.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        style_run(r, 14.5, COL["white"] if band == COL["primary"] else COL["text"], bold=True)
        add_bullets(s, x + 0.32, cy + 0.85, cw - 0.6, 3.3, bullets,
                    size=SZ["card_bullet"], marker=COL["secondary"], space_after=8)
    pic_cover(s, d["image"], x2, cy, 3.0, ch, radius=0.08)
    note_sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 6.3, 12.26, 0.55,
                    fill=COL["accent_light"], radius=0.3)
    tf = note_sp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = d["split_note"]
    style_run(r, 14, COL["primary_dark"], bold=True)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_examples(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    n = len(d["cards"])
    gap = LAY["gap_in"]
    m = LAY["margin_in"]
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 1.85, 4.4
    for i, c in enumerate(d["cards"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch, fill=COL["soft_bg"])
        txt(s, x + 0.2, cy + 0.45, cw - 0.4, 0.8, c["stat"], size=34,
            color=COL["primary"], bold=True, align=PP_ALIGN.CENTER, font=FONT_H)
        txt(s, x + 0.25, cy + 1.4, cw - 0.5, 0.75, c["title"], size=15.5,
            color=COL["primary_dark"], bold=True, align=PP_ALIGN.CENTER,
            line_spacing=1.1)
        txt(s, x + 0.28, cy + 2.35, cw - 0.56, 1.8, c["label"], size=13,
            color=COL["text_muted"], align=PP_ALIGN.CENTER, line_spacing=1.18)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_examples_with_banner(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    pic_cover(s, d["image"], 0.7, 1.68, 11.93, 2.18, radius=0.07)
    n = len(d["cards"])
    gap = 0.3
    m = LAY["margin_in"]
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 4.12, 2.42
    for i, c in enumerate(d["cards"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch, fill=COL["soft_bg"])
        txt(s, x + 0.2, cy + 0.2, cw - 0.4, 0.55, c["stat"], size=26,
            color=COL["primary"], bold=True, align=PP_ALIGN.CENTER, font=FONT_H)
        txt(s, x + 0.22, cy + 0.8, cw - 0.44, 0.55, c["title"], size=13,
            color=COL["primary_dark"], bold=True, align=PP_ALIGN.CENTER,
            line_spacing=1.05)
        txt(s, x + 0.24, cy + 1.38, cw - 0.48, 0.95, c["label"], size=11.5,
            color=COL["text_muted"], align=PP_ALIGN.CENTER, line_spacing=1.12)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_timeline(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    n = len(d["phases"])
    m = 0.6
    gap = 0.34
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 1.85, 4.75
    for i, ph in enumerate(d["phases"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch)
        chip(s, x + 0.18, cy + 0.2, cw - 0.36, 0.5, ph["period"], 14.5,
             COL["primary"], COL["white"])
        txt(s, x + 0.15, cy + 0.85, cw - 0.3, 0.45, ph["title"], size=18,
            color=COL["primary_dark"], bold=True, align=PP_ALIGN.CENTER, font=FONT_H)
        add_bullets(s, x + 0.28, cy + 1.5, cw - 0.5, 3.0, ph["items"],
                    size=SZ["timeline_item"], marker=COL["secondary"], space_after=8)
        if i < n - 1:
            ax = x + cw + gap / 2 - 0.14
            shape(s, MSO_SHAPE.RIGHT_ARROW, ax, cy + ch / 2 - 0.15, 0.28, 0.3,
                  fill=COL["secondary"])
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_table(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    tdata = d["table"]
    rows = 1 + len(tdata["rows"])
    cols = len(tdata["headers"])
    x, y, w = 0.7, 1.7, 11.93
    gr = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(0.4 * rows))
    table = gr.table
    widths = [0.55, 2.85, 5.73, 2.8]
    for i, cw in enumerate(widths):
        table.columns[i].width = Inches(cw)
    table.rows[0].height = Inches(0.42)
    for ri in range(1, rows):
        table.rows[ri].height = Inches(0.42)
    aligns = [PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT]
    for ci, htext in enumerate(tdata["headers"]):
        cell = table.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COL["primary"]
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left, cell.margin_right = Inches(0.08), Inches(0.08)
        tf = cell.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = aligns[ci]
        r = p.add_run()
        r.text = htext
        style_run(r, SZ["table_header"], COL["white"], bold=True)
    for ri, row in enumerate(tdata["rows"], start=1):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COL["table_alt"] if ri % 2 == 0 else COL["bg"]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left, cell.margin_right = Inches(0.08), Inches(0.08)
            tf = cell.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = aligns[ci]
            r = p.add_run()
            r.text = val
            style_run(r, SZ["table_body"], COL["text"],
                      bold=(ci == 1))
    ty = y + 0.42 * rows + 0.22
    bar = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, ty, 11.93, 0.62,
                fill=COL["primary_dark"], radius=0.12)
    txt(s, 1.0, ty + 0.08, 5.5, 0.46, tdata["total_label"], size=16,
        color=COL["white"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 6.8, ty + 0.06, 5.55, 0.5, tdata["total_value"], size=19,
        color=COL["accent"], bold=True, align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE, font=FONT_H)
    txt(s, 0.7, ty + 0.75, 11.93, 0.6, tdata["total_note"], size=13,
        color=COL["text_muted"], align=PP_ALIGN.CENTER, line_spacing=1.1)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_sponsors(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    n = len(d["cards"])
    gap = LAY["gap_in"]
    m = LAY["margin_in"]
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 1.8, 4.25
    for i, c in enumerate(d["cards"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch)
        band = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, cy, cw, 0.6,
                     fill=COL["primary"], radius=0.12)
        tf = band.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = c["title"]
        style_run(r, 16.5, COL["white"], bold=True)
        txt(s, x + 0.2, cy + 0.82, cw - 0.4, 0.6, c["price"], size=26,
            color=COL["primary_dark"], bold=True, align=PP_ALIGN.CENTER, font=FONT_H)
        add_bullets(s, x + 0.32, cy + 1.7, cw - 0.6, 2.3, c["bullets"],
                    size=SZ["card_bullet"], marker=COL["secondary"], space_after=8)
    bar = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, m, 6.3, SW - 2 * m, 0.55,
                fill=COL["accent_light"], radius=0.3)
    tf = bar.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = d["note_line"]
    style_run(r, 14, COL["primary_dark"], bold=True)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_benefits(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    n = len(d["columns"])
    gap = 0.5
    m = 1.0
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 1.8, 4.75
    bands = [COL["primary"], COL["accent"]]
    for i, c in enumerate(d["columns"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch)
        band = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, cy, cw, 0.68,
                     fill=bands[i], radius=0.1)
        tf = band.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = c["title"]
        style_run(r, 19, COL["white"] if i == 0 else COL["text"], bold=True)
        add_bullets(s, x + 0.45, cy + 1.05, cw - 0.85, 3.4, c["items"],
                    size=15.5, marker=COL["secondary"], space_after=13)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_risks(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    n = len(d["cards"])
    gap = LAY["gap_in"]
    m = LAY["margin_in"]
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 1.95, 4.2
    for i, c in enumerate(d["cards"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch, fill=COL["soft_bg"])
        txt(s, x + 0.25, cy + 0.4, cw - 0.5, 0.85, c["title"], size=18,
            color=COL["primary_dark"], bold=True, align=PP_ALIGN.CENTER,
            line_spacing=1.1)
        chip(s, x + (cw - 1.3) / 2, cy + 1.5, 1.3, 0.42, "SOLUSI", 12,
             COL["primary"], COL["white"])
        txt(s, x + 0.3, cy + 2.15, cw - 0.6, 1.8, c["fix"], size=13.5,
            color=COL["text_muted"], align=PP_ALIGN.CENTER, line_spacing=1.2)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_indicators(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"])
    n = len(d["cards"])
    gap = 0.3
    m = LAY["margin_in"]
    cw = (SW - 2 * m - (n - 1) * gap) / n
    cy, ch = 1.75, 2.35
    for i, c in enumerate(d["cards"]):
        x = m + i * (cw + gap)
        card(s, x, cy, cw, ch, fill=COL["soft_bg"])
        txt(s, x + 0.15, cy + 0.3, cw - 0.3, 0.65, c["stat"], size=30,
            color=COL["primary"], bold=True, align=PP_ALIGN.CENTER, font=FONT_H)
        txt(s, x + 0.2, cy + 1.1, cw - 0.4, 1.1, c["label"], size=12.5,
            color=COL["text_muted"], align=PP_ALIGN.CENTER, line_spacing=1.15)
    by = 4.5
    for b in d["bullets"]:
        chip(s, m, by, SW - 2 * m, 0.6, b, 15, COL["primary_light"],
             COL["primary_dark"], align=PP_ALIGN.LEFT)
        by += 0.76
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_closing(prs, d, page):
    s = blank_slide(prs)
    txt(s, 0.7, 1.55, 6.9, 0.8, d["headline"], size=34,
        color=COL["primary_dark"], bold=True, font=FONT_H)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 2.45, 1.4, 0.07,
          fill=COL["accent"], radius=0.5)
    by = 2.85
    for b in d["cta_bullets"]:
        chip(s, 0.7, by, 6.9, 0.66, b, 15, COL["primary_light"],
             COL["primary_dark"], align=PP_ALIGN.LEFT)
        by += 0.82
    chip(s, 0.7, by + 0.1, 6.9, 0.62, d["closing_line"], 14.5,
         COL["accent_light"], COL["primary_dark"])
    txt(s, 0.7, by + 0.95, 6.9, 0.5, "Kontak: " + d["contact"], size=12.5,
        color=COL["text_muted"], line_spacing=1.15)
    pic_cover(s, d["image"], 7.95, 1.6, 4.7, 5.0, radius=0.09)
    footer(s, page)
    notes(s, d["notes"])
    return s


def layout_mechanism2(prs, d, page):
    s = blank_slide(prs)
    header(s, d["headline"], kicker=d.get("kicker"), subhead=d.get("subhead"))

    m = LAY["margin_in"]
    arrow_gap = 0.40
    steps = d["steps"]
    n = len(steps)
    cw = (SW - 2 * m - (n - 1) * arrow_gap) / n
    cy, ch = 1.92, 2.48

    for i, st in enumerate(steps):
        x = m + i * (cw + arrow_gap)
        card(s, x, cy, cw, ch, fill=COL["soft_bg"])
        photos = st["photos"]
        ph = 1.14
        if len(photos) == 1:
            pic_cover(s, photos[0], x + 0.09, cy + 0.08, cw - 0.18, ph, radius=0.05)
        else:
            half = (cw - 0.18 - 0.06) / 2
            for j, key in enumerate(photos):
                pic_cover(s, key, x + 0.09 + j * (half + 0.06), cy + 0.08, half, ph,
                          radius=0.05)
        badge = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.15, cy + 0.15,
                      0.40, 0.34, fill=COL["accent"], radius=0.5)
        btf = badge.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run()
        br.text = str(i + 1)
        style_run(br, 13, COL["white"], bold=True)
        txt(s, x + 0.10, cy + 1.28, cw - 0.20, 0.30, st["title"], size=13,
            color=COL["primary_dark"], bold=True, align=PP_ALIGN.CENTER,
            font=FONT_H)
        txt(s, x + 0.13, cy + 1.60, cw - 0.26, 0.82, st["caption"], size=10,
            color=COL["text_muted"], align=PP_ALIGN.CENTER, line_spacing=1.12)
        if i < n - 1:
            tri = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, x + cw + 0.09,
                        cy + 0.50, 0.22, 0.24, fill=COL["secondary"])
            tri.rotation = 90

    loop_y = 4.50
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, m + 0.26, loop_y, SW - 2 * m - 0.26,
          0.34, fill=COL["primary_light"], radius=0.5)
    txt(s, m + 0.40, loop_y + 0.06, SW - 2 * m - 0.68, 0.24, d["loop_text"],
        size=11, color=COL["primary_dark"], bold=True, align=PP_ALIGN.CENTER)
    lt = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, m + 0.03, loop_y + 0.08,
               0.20, 0.19, fill=COL["primary"])
    lt.rotation = 270

    by, bh = 5.00, 1.46
    fx, fw = m, 6.00
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, fx, by, fw, bh,
          fill=COL["accent_light"], radius=0.06)
    pic_cover(s, d["frass"]["image"], fx + 0.16, by + 0.14, 1.98, bh - 0.28,
              radius=0.05)
    tx = fx + 2.26
    tw = fw - 2.26 - 0.16
    txt(s, tx, by + 0.18, tw, 0.28, d["frass"]["title"], size=12.5,
        color=COL["primary_dark"], bold=True, font=FONT_H)
    txt(s, tx, by + 0.52, tw, 0.80, d["frass"]["body"], size=11,
        color=COL["text"], line_spacing=1.16)

    sx = fx + fw + 0.22
    sw = SW - m - sx
    card(s, sx, by, sw, bh, fill=COL["bg"])
    txt(s, sx + 0.24, by + 0.16, sw - 0.48, 0.26, d["facts"]["title"],
        size=12.5, color=COL["primary_dark"], bold=True, font=FONT_H)
    ry = by + 0.50
    for f in d["facts"]["rows"]:
        dot = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, sx + 0.26, ry + 0.07,
                    0.13, 0.13, fill=COL["secondary"], radius=0.5)
        dot.line.fill.background()
        txt(s, sx + 0.52, ry, sw - 0.78, 0.28, f, size=11,
            color=COL["text"], line_spacing=1.05)
        ry += 0.30

    txt(s, m, 6.58, SW - 2 * m, 0.26, d["source_line"], size=8.5,
        color=COL["text_muted"], align=PP_ALIGN.CENTER)
    footer(s, page)
    notes(s, d["notes"])
    return s


def build(out_path=None):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    for d in OUTLINE["slides"]:
        layout = d["layout"]
        page = d["id"]
        if layout == "cover":
            layout_cover(prs, d)
        elif layout == "exec_summary":
            layout_exec_summary(prs, d)
        elif layout == "content_stats":
            layout_content_stats(prs, d, page)
        elif layout == "content_image":
            layout_content_image(prs, d, page)
        elif layout == "pillars":
            layout_pillars(prs, d, page)
        elif layout == "diagram":
            layout_diagram(prs, d, page)
        elif layout == "photo_captions":
            layout_photo_captions(prs, d, page)
        elif layout == "two_column":
            layout_two_column(prs, d, page)
        elif layout == "examples":
            layout_examples(prs, d, page)
        elif layout == "examples_with_banner":
            layout_examples_with_banner(prs, d, page)
        elif layout == "timeline":
            layout_timeline(prs, d, page)
        elif layout == "table":
            layout_table(prs, d, page)
        elif layout == "sponsors":
            layout_sponsors(prs, d, page)
        elif layout == "benefits":
            layout_benefits(prs, d, page)
        elif layout == "risks":
            layout_risks(prs, d, page)
        elif layout == "indicators":
            layout_indicators(prs, d, page)
        elif layout == "closing":
            layout_closing(prs, d, page)
        elif layout == "mechanism2":
            layout_mechanism2(prs, d, page)
        else:
            raise ValueError(f"unknown layout: {layout}")
    out = pathlib.Path(out_path or OUT_PATH)
    prs.save(str(out))
    print(f"Saved {out} with {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--outline", default=str(OUTLINE_PATH))
    ap.add_argument("--out", default=str(OUT_PATH))
    args = ap.parse_args()
    load_outline(args.outline)
    build(args.out)
